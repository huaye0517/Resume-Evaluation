# -*- coding: utf-8 -*-
"""Flask 应用入口。

提供三个接口：
- GET  /            首页（上传/粘贴界面）
- POST /evaluate    接收 JD + 多份简历，逐份评估并返回 JSON
- POST /export      接收评估结果，生成 Excel 供下载
"""

import io
import os
from datetime import datetime

from dotenv import load_dotenv
from flask import Flask, jsonify, render_template, request, send_file

import resume_parser
from evaluator import build_parse_error_result, evaluate_resume

# 启动时加载 .env 中的配置
load_dotenv()

app = Flask(__name__)
# 限制上传大小，避免超大文件拖垮服务（单次合计 50MB）
app.config["MAX_CONTENT_LENGTH"] = 50 * 1024 * 1024


@app.route("/")
def index():
    """首页。"""
    return render_template("index.html")


def _resolve_jd_text():
    """从请求中解析 JD 文本：优先用上传的 JD 文件，否则用粘贴的文本。"""
    jd_file = request.files.get("jd_file")
    if jd_file and jd_file.filename:
        if not resume_parser.is_supported(jd_file.filename):
            return None, f"JD 文件格式不支持：{jd_file.filename}"
        text, error = resume_parser.parse_file(jd_file.filename, jd_file.read())
        if error:
            return None, f"JD 文件解析失败：{error}"
        return text, None

    jd_text = (request.form.get("jd_text") or "").strip()
    if jd_text:
        return jd_text, None

    return None, "请粘贴 JD 文本或上传 JD 文件。"


@app.route("/evaluate", methods=["POST"])
def evaluate():
    """接收 JD 与多份简历，逐份评估并返回排序后的结果。"""
    jd_text, jd_error = _resolve_jd_text()
    if jd_error:
        return jsonify({"ok": False, "message": jd_error}), 400

    resume_files = request.files.getlist("resume_files")
    resume_files = [f for f in resume_files if f and f.filename]
    if not resume_files:
        return jsonify({"ok": False, "message": "请至少上传一份简历。"}), 400

    results = []
    for f in resume_files:
        filename = f.filename
        if not resume_parser.is_supported(filename):
            results.append(build_parse_error_result(filename, "不支持的文件格式（仅支持 PDF / Word(.docx) / txt）"))
            continue

        text, parse_error = resume_parser.parse_file(filename, f.read())
        if parse_error:
            results.append(build_parse_error_result(filename, parse_error))
            continue

        results.append(evaluate_resume(jd_text, text, filename))

    # 按匹配度从高到低排序，便于优先查看
    results.sort(key=lambda r: r.get("score", 0), reverse=True)

    return jsonify({"ok": True, "results": results})


@app.route("/export", methods=["POST"])
def export():
    """把前端传回的评估结果导出为 Excel 文件供下载。"""
    payload = request.get_json(silent=True) or {}
    results = payload.get("results") or []
    if not results:
        return jsonify({"ok": False, "message": "没有可导出的结果。"}), 400

    workbook = _build_workbook(results)

    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)

    filename = f"简历评估汇总_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return send_file(
        buffer,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def _build_workbook(results):
    """根据评估结果列表构建 openpyxl 工作簿。"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "评估汇总"

    headers = [
        "文件名", "候选人", "结论", "匹配度",
        "岗位方向匹配", "相关经验", "专业技能", "稳定性与成长性",
        "亮点", "风险点", "建议面试问题", "结论说明",
    ]
    sheet.append(headers)

    # 表头样式
    header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for r in results:
        dims = r.get("dimension_scores") or {}
        sheet.append([
            r.get("filename", ""),
            r.get("candidate_name", ""),
            r.get("verdict", ""),
            r.get("score", 0),
            dims.get("岗位方向匹配", ""),
            dims.get("相关经验", ""),
            dims.get("专业技能", ""),
            dims.get("稳定性与成长性", ""),
            "\n".join(r.get("highlights") or []),
            "\n".join(r.get("risks") or []),
            "\n".join(r.get("interview_questions") or []),
            r.get("summary", ""),
        ])

    # 列宽与自动换行设置
    column_widths = [22, 12, 8, 8, 12, 10, 10, 14, 36, 36, 40, 40]
    for idx, width in enumerate(column_widths, start=1):
        sheet.column_dimensions[chr(64 + idx)].width = width

    wrap_alignment = Alignment(wrap_text=True, vertical="top")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = wrap_alignment

    return workbook


if __name__ == "__main__":
    # 启动时提示是否检测到 API Key，便于排查配置问题
    # 注意：修改 .env 后必须重启本服务才会生效（load_dotenv 仅在启动时读取一次）
    if os.environ.get("DEEPSEEK_API_KEY", "").strip():
        print("[配置检查] 已检测到 DEEPSEEK_API_KEY，可以正常评估。")
    else:
        print("[配置检查] 未检测到 DEEPSEEK_API_KEY，请在 .env 中填写后重启本服务。")

    # 本地运行，浏览器访问 http://127.0.0.1:5000
    app.run(host="127.0.0.1", port=5000, debug=True)
